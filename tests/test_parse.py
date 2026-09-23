from pathlib import Path

import pytest

from src.parse import SUPPORTED_SUFFIXES, parse_document, read_text, split_clauses


def test_inline_clauses_and_sections():
    text = (
        "Преамбула\n2. Общие положения\n2.4.7. Правила.\n"
        "3. Работники\n3.4. Соблюдение распорядка. 3.10.Работники обязаны.\n"
        "5. Ограничения\n5.8. Запреты\n5.8.1. Участие запрещено."
    )
    fragments = split_clauses(text, "before/rules.txt")
    assert len(fragments) == 6
    assert [f["clause"] for f in fragments] == ["", "2.4.7", "3.4", "3.10", "5.8", "5.8.1"]
    assert fragments[2] == {
        "doc": "before/rules.txt", "clause": "3.4",
        "text": "3.4. Соблюдение распорядка.", "section": "3. Работники",
    }
    assert fragments[3]["text"] == "3.10.Работники обязаны."
    assert fragments[-1]["section"] == "5. Ограничения"


def test_dates_and_references_are_not_boundaries():
    text = "3.4. Утверждено 25.06.2021. См. п. 5.8. и пункт 2.4.7. документа."
    assert split_clauses(text) == [
        {"doc": "", "clause": "3.4", "text": text, "section": ""}
    ]


def test_unnumbered_and_empty():
    assert split_clauses(" \n") == []
    assert split_clauses("Обычный текст")[0]["clause"] == ""
    assert split_clauses("3.4.Текст\nпродолжение")[0]["text"].endswith("продолжение")


def test_inline_section_heading():
    fragments = split_clauses("9.1. Конец пункта. 10.Контроль качества\n10.1. Проверка")
    assert len(fragments) == 2
    assert fragments[0]["text"] == "9.1. Конец пункта."
    assert fragments[1]["section"] == "10.Контроль качества"


@pytest.mark.parametrize("encoding", ["utf-8-sig", "cp1251"])
def test_txt(tmp_path, encoding):
    path = tmp_path / "document.TXT"
    path.write_bytes("3.4. Работники\n5.8. Запреты".encode(encoding))
    assert [f["clause"] for f in parse_document(path)] == ["3.4", "5.8"]


def test_docx_tables_in_document_order(tmp_path):
    from docx import Document

    document = Document()
    document.add_paragraph("3. Работники")
    document.add_paragraph("3.4. Первый пункт")
    table = document.add_table(rows=1, cols=2)
    cell = table.cell(0, 0).merge(table.cell(0, 1))
    cell.text = "3.5. Пункт в таблице"
    nested = cell.add_table(rows=1, cols=1)
    nested.cell(0, 0).text = "3.6. Вложенная таблица"
    document.add_paragraph("5. Ограничения")
    document.add_paragraph("5.8. Последний пункт")
    path = tmp_path / "document.docx"
    document.save(path)
    fragments = parse_document(path)
    assert [f["clause"] for f in fragments] == ["3.4", "3.5", "3.6", "5.8"]
    assert fragments[2]["section"] == "3. Работники"


def test_pdf_all_pages(tmp_path):
    import pymupdf

    path = tmp_path / "document.pdf"
    with pymupdf.open() as document:
        document.new_page().insert_text((72, 72), "3. Employees\n3.4. First clause")
        document.new_page().insert_text((72, 72), "5. Restrictions\n5.8. Last clause")
        document.save(path)
    fragments = parse_document(path)
    assert [f["clause"] for f in fragments] == ["3.4", "5.8"]
    assert fragments[-1]["section"] == "5. Restrictions"


def test_xlsx_all_sheets(tmp_path):
    from openpyxl import Workbook

    workbook = Workbook()
    workbook.active.append(["3. Работники"])
    workbook.active.append(["3.4.", "Первый пункт", 0, False])
    sheet = workbook.create_sheet("После")
    sheet.append(["5. Ограничения"])
    sheet.append(["5.8.", "Последний пункт"])
    path = tmp_path / "document.xlsx"
    workbook.save(path)
    workbook.close()
    fragments = parse_document(path)
    assert [f["clause"] for f in fragments] == ["3.4", "5.8"]
    assert "Первый пункт\t0\tFalse" in fragments[0]["text"]


def test_unsupported_format(tmp_path):
    with pytest.raises(ValueError, match="Неподдерживаемый формат"):
        read_text(tmp_path / "document.csv")


@pytest.mark.parametrize("edition", ["before", "after"])
def test_sample_editions(edition):
    root = Path(__file__).resolve().parents[1] / "data" / "sample" / edition
    paths = sorted(p for p in root.rglob("*") if p.suffix.lower() in SUPPORTED_SUFFIXES)
    if not paths:
        pytest.skip(f"Отсутствуют исходные документы: {root}")
    clauses = set()
    expected_counts = {
        "Положение_о_внутреннем_аудите_редакция_8_обезличено.docx": 320,
        "Положение_о_внутреннем_аудите_редакция_9_обезличено.docx": 309,
    }
    for path in paths:
        fragments = parse_document(path)
        assert len(fragments) > 0, f"Нет фрагментов в {path}"
        if path.name in expected_counts:
            assert len(fragments) == expected_counts[path.name]
            by_clause = {f["clause"]: f for f in fragments}
            for number in (10, 11, 12, 13):
                assert by_clause[f"{number}.1"]["section"].startswith(f"{number}.")
        print(f"{path}: {len(fragments)} фрагментов")
        assert all(set(f) == {"doc", "clause", "text", "section"} for f in fragments)
        clauses.update(f["clause"] for f in fragments)
    assert {"3.4", "5.8"} <= clauses, f"В редакции {edition} отсутствуют обязательные пункты"
