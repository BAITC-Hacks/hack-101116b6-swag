"""Локальное извлечение текста и пунктов документов без API и GPU.

Запуск: python -m src.parse data/sample --limit 10
"""

import argparse
import json
from pathlib import Path
import re
from typing import TypedDict
import warnings


SUPPORTED_SUFFIXES = {".docx", ".pdf", ".xlsx", ".txt"}


class Fragment(TypedDict):
    doc: str
    clause: str
    text: str
    section: str


def read_text(path: str | Path) -> str:
    """Прочитать документ; TXT — UTF-8 (включая BOM), затем Windows-1251.

    PDF должен содержать текстовый слой: OCR не выполняется. В XLSX читаются
    все листы, формулы сохраняются как формулы, пустые ячейки — как столбцы.
    Автонумерация Word не восстанавливается: номера должны быть в тексте.
    PDF без извлекаемого текста вызывает ValueError; пропущенные или пустые
    страницы в частично прочитанном PDF сопровождаются UserWarning.
    """
    path = Path(path)
    suffix = path.suffix.lower()
    if suffix == ".txt":
        raw = path.read_bytes()
        try:
            return raw.decode("utf-8-sig")
        except UnicodeDecodeError:
            return raw.decode("cp1251")
    if suffix == ".docx":
        from docx import Document
        from docx.table import Table

        def blocks(container):
            for block in container.iter_inner_content():
                if isinstance(block, Table):
                    seen = set()
                    for row in block.rows:
                        for cell in row.cells:
                            # Объединённая ячейка может встречаться несколько раз.
                            if cell._tc not in seen:
                                seen.add(cell._tc)
                                yield from blocks(cell)
                else:
                    yield block.text

        return "\n".join(blocks(Document(path)))
    if suffix == ".pdf":
        import pymupdf

        with pymupdf.open(path) as document:
            pages = []
            incomplete = []
            for index in range(len(document)):
                try:
                    page_text = document[index].get_text(sort=True)
                except Exception as exc:
                    # Ошибка одной страницы не должна скрывать остальные.
                    pages.append("")
                    incomplete.append(f"{index + 1} (ошибка {type(exc).__name__})")
                    continue
                pages.append(page_text)
                if not page_text.strip():
                    incomplete.append(f"{index + 1} (нет текста)")
            text = "\n".join(pages)
            if not text.strip():
                raise ValueError(
                    f"PDF {path}: нет извлекаемого текста. "
                    "OCR не выполняется; нужен документ с доступным текстовым слоем."
                )
            if incomplete:
                warnings.warn(
                    f"PDF {path}: возможно неполное извлечение текста; "
                    f"страницы: {', '.join(incomplete)}. "
                    "Пустая страница не обязательно является сканом. OCR не выполняется.",
                    UserWarning,
                    stacklevel=2,
                )
            return text
    if suffix == ".xlsx":
        from openpyxl import load_workbook

        workbook = load_workbook(path, read_only=True, data_only=False)
        try:
            return "\n".join(
                "\t".join("" if value is None else str(value) for value in row)
                for sheet in workbook
                for row in sheet.iter_rows(values_only=True)
            )
        finally:
            workbook.close()
    raise ValueError(f"Неподдерживаемый формат: {suffix or path.name}")


# Не распознаём даты и части чисел как номера пунктов.
_CLAUSE = re.compile(r"(?<![\w.])(?P<number>\d+(?:\.\d+)+)\.(?!\d)")
_DATE = re.compile(r"\d{1,2}\.\d{1,2}\.\d{4}$")
_REFERENCE = re.compile(r"(?:\bп\.|\bпп\.|\bпункт[а-я]*|\bподпункт[а-я]*)[ \t]*\Z", re.I)
_SECTION = re.compile(
    r"(?:^[ \t]*|(?<=[.;])[ \t]+)(?:Раздел[ \t]+)?(?P<number>\d+)[.)][ \t]*"
    r"(?P<title>[^\W\d_][^\n]*)$", re.M | re.I
)


def split_clauses(text: str, doc: str = "") -> list[Fragment]:
    """Разбить текст на пункты, включая номера посреди строки.

    clause хранится без конечной точки, text включает исходный номер.
    section — ближайший предшествующий заголовок вида «3. Заголовок».
    Преамбула и ненумерованный текст сохраняются с clause="";
    при отсутствии заголовка section="". Заголовки не входят в текст пункта.
    """
    text = text.replace("\r\n", "\n").replace("\r", "\n")
    events = [(m.start(), m.end(), "section", m) for m in _SECTION.finditer(text)]
    for match in _CLAUSE.finditer(text):
        if _DATE.fullmatch(match["number"]):
            continue
        if _REFERENCE.search(text[max(0, match.start() - 40):match.start()]):
            continue
        # Заголовок не должен поглощать пункт на той же строке.
        events.append((match.start(), match.end(), "clause", match))
    events.sort(key=lambda event: event[0])
    fragments: list[Fragment] = []
    start, clause, section = 0, "", ""

    def append(end: int) -> None:
        body = text[start:end].strip()
        if body:
            fragments.append({"doc": doc, "clause": clause, "text": body, "section": section})

    for index, (position, end, kind, match) in enumerate(events):
        append(position)
        if kind == "section":
            if index + 1 < len(events):
                end = min(end, events[index + 1][0])
            section = text[position:end].strip()
            start, clause = end, ""
        else:
            start, clause = position, match["number"]
    append(len(text))
    return fragments


def parse_document(path: str | Path) -> list[Fragment]:
    """Путь в doc сохраняет различие редакций с одинаковыми именами файлов."""
    return split_clauses(read_text(path), doc=str(path))


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("path", type=Path, help="Документ или каталог документов")
    parser.add_argument("--limit", type=int, default=10)
    args = parser.parse_args()
    if args.limit < 0:
        parser.error("--limit должен быть неотрицательным")
    paths = [args.path] if args.path.is_file() else sorted(
        p for p in args.path.rglob("*") if p.is_file() and p.suffix.lower() in SUPPORTED_SUFFIXES
    )
    if not paths:
        parser.error(f"Документы не найдены: {args.path}")
    for path in paths:
        fragments = parse_document(path)
        print(f"{path}: {len(fragments)} фрагментов")
        print(json.dumps(fragments[:args.limit], ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
