"""Portable report containing only entries with verified source quotations."""
from io import BytesIO

from docx import Document


def build_docx(result):
    document = Document()
    document.add_heading("Сравнение документов", 0)
    document.add_paragraph("Выводы рекомендательные и требуют проверки ответственным человеком. "
                           "В отчёт включены только записи с подтверждёнными цитатами.")
    count = 0
    for group, title in (("findings", "Находки"), ("units", "Подразделения"),
                         ("function_map", "Сопоставление функций")):
        document.add_heading(title, 1)
        for item in result.get(group, []):
            evidence = item.get("evidence") or []
            if not evidence or not all(e.get("verified") is True and e.get("quote")
                                       and e.get("doc") and e.get("clause") for e in evidence):
                continue
            count += 1
            document.add_heading(item.get("title") or item.get("name") or "Сопоставление", 2)
            if group == "function_map":
                document.add_paragraph(f"До: {item.get('before_unit') or '—'} · {item.get('before_function') or '—'}")
                document.add_paragraph(f"После: {item.get('after_unit') or '—'} · {item.get('after_function') or '—'}")
            else:
                document.add_paragraph(item.get("description") or "Сведения извлечены из документов; см. источники.")
            if item.get("status"):
                document.add_paragraph(f"Статус сопоставления: {item['status']}")
            for source in evidence:
                document.add_paragraph(f"{source['doc']} · пункт {source['clause']}")
                document.add_paragraph(source["quote"])
    if not count:
        document.add_paragraph("Нет записей с подтверждёнными источниками для включения в отчёт.")
    stream = BytesIO()
    document.save(stream)
    return stream.getvalue()
